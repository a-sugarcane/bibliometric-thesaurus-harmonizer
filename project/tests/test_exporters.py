"""Unit tests for dual-metric audit exporter and VOSviewer export."""

import tempfile
import unittest
from pathlib import Path
from core.parsers.base_parser import DocumentRecord
from core.engine.tier_classifier import HarmonizationRule
from core.exporters.audit_reporter import AuditReporter
from core.exporters.vosviewer import export_vosviewer_thesaurus


class TestExporters(unittest.TestCase):

    def test_dual_metric_frequency_calculation(self):
        docs = [
            DocumentRecord(
                doc_id="WOS:001",
                effective_keywords={"prostate cancer", "prostate neoplasms"},
                author_keywords_de={"prostate cancer", "prostate neoplasms"},
            ),
            DocumentRecord(
                doc_id="WOS:002",
                effective_keywords={"prostate cancer"},
                author_keywords_de={"prostate cancer"},
            ),
        ]
        mapping = {"prostate neoplasms": "prostate cancer", "prostate cancer": "prostate cancer"}
        reporter = AuditReporter()
        metrics = reporter.calculate_frequencies(docs, mapping)

        # Canonical term: prostate cancer
        # Doc 1 has 2 occurrences mapped to "prostate cancer"
        # Doc 2 has 1 occurrence mapped to "prostate cancer"
        # Raw sum = 3, Independent doc count = 2, Deduplication reduction = 1
        self.assertIn("prostate cancer", metrics)
        self.assertEqual(metrics["prostate cancer"]["raw_sum"], 3)
        self.assertEqual(metrics["prostate cancer"]["doc_count"], 2)
        self.assertEqual(metrics["prostate cancer"]["reduction"], 1)

    def test_vosviewer_thesaurus_export(self):
        rules = [
            HarmonizationRule(
                raw_term="prostate cancers",
                target_term="prostate cancer",
                tier="Tier 1",
                rule_source="Lemmatizer",
                clinical_safety_check="Passed",
                selected_for_export=True,
            ),
            # Self-loop should be omitted
            HarmonizationRule(
                raw_term="prostate cancer",
                target_term="prostate cancer",
                tier="Tier 1",
                rule_source="Identity",
                clinical_safety_check="Passed",
                selected_for_export=True,
            ),
            # Unselected rule should be omitted by default
            HarmonizationRule(
                raw_term="mCRPC",
                target_term="prostate cancer",
                tier="Tier 3",
                rule_source="High Risk",
                clinical_safety_check="Blocked",
                selected_for_export=False,
            ),
        ]

        with tempfile.NamedTemporaryFile("w+", delete=False, suffix=".txt") as tmp:
            tmp_path = Path(tmp.name)

        out_path = export_vosviewer_thesaurus(rules, tmp_path)
        with open(out_path, "r", encoding="utf-8") as f:
            lines = [line.strip() for line in f if line.strip()]

        self.assertEqual(lines[0], "label\treplace by")
        self.assertEqual(lines[1], "prostate cancers\tprostate cancer")
        self.assertEqual(len(lines), 2)  # Header + 1 rule only

    def test_table_s1_excel_export(self):
        rules = [
            HarmonizationRule(
                raw_term="prostate cancers",
                target_term="prostate cancer",
                tier="Tier 1",
                rule_source="Lemmatizer",
                clinical_safety_check="Passed",
                raw_frequency=15,
                selected_for_export=True,
            )
        ]
        docs = [
            DocumentRecord(
                doc_id="WOS:001",
                effective_keywords={"prostate cancers"},
            )
        ]
        with tempfile.NamedTemporaryFile("w+", delete=False, suffix=".xlsx") as tmp:
            tmp_path = Path(tmp.name)

        reporter = AuditReporter()
        out_path = reporter.export_table_s1(rules, docs, tmp_path)
        self.assertTrue(out_path.exists())

        import openpyxl
        wb = openpyxl.load_workbook(out_path)
        self.assertIn("Table S1 Harmonization Audit", wb.sheetnames)
        self.assertIn("Methodological Statement", wb.sheetnames)

        ws1 = wb["Table S1 Harmonization Audit"]
        # Row 1 is header, Row 2 is data
        self.assertEqual(ws1.cell(row=2, column=1).value, "prostate cancers")
        self.assertEqual(ws1.cell(row=2, column=2).value, "prostate cancer")
        self.assertEqual(ws1.cell(row=2, column=7).value, 1)  # Independent Doc Count



if __name__ == "__main__":
    unittest.main()
